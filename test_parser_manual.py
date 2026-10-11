
from pathlib import Path
from docx import Document

from app.services.resume_parser import extract_resume_text


# Create a sample DOCX resume for testing
file_path = Path("sample_resume.docx")

document = Document()
document.add_heading("Sample Python Developer Resume", level=1)
document.add_paragraph(
    "Skills: Python, Django, FastAPI, SQL, PostgreSQL"
)
document.add_paragraph(
    "Projects: Customer Sales Analytics using Python and SQL"
)
document.save(file_path)

# Extract the text
text = extract_resume_text(str(file_path))

print("\n--- EXTRACTED RESUME TEXT ---\n")
print(text)

assert "Python" in text
assert "Django" in text
assert "Customer Sales Analytics" in text

print("\nSUCCESS: Resume text extraction is working!")